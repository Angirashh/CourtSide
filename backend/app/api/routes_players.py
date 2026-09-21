import csv
import io
import uuid
from datetime import datetime
from typing import List, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session

from app.api.dependencies import CurrentUser, ensure_tournament_access, get_db, require_operator_or_organiser, require_organiser
from app.db import models
from app.schemas.player import (
    PlayerCreate,
    PlayerBatchCreate,
    PlayerResponse,
    PlayerWithdrawRequest,
    PlayerWithdrawResponse,
)
from app.services.match_progression import resolve_pending_walkovers

router = APIRouter(prefix="/tournaments/{tournament_id}/players", tags=["Players"])


@router.post("", response_model=PlayerResponse, status_code=status.HTTP_201_CREATED)
def register_player(
    tournament_id: str,
    payload: PlayerCreate,
    current_user: CurrentUser = Depends(require_organiser),
    db: Session = Depends(get_db)
):
    """Registers a single player into a tournament."""
    tournament = db.query(models.Tournament).filter(models.Tournament.id == tournament_id).first()
    if not tournament:
        raise HTTPException(status_code=404, detail="Tournament not found")
    ensure_tournament_access(tournament, current_user, db)

    if tournament.status != models.TournamentStatus.DRAFT:
        raise HTTPException(
            status_code=400,
            detail="Cannot add players once scheduling has started or the tournament is in progress."
        )

    player = models.Player(
        id=f"P_{uuid.uuid4().hex[:8]}",
        tournament_id=tournament_id,
        name=payload.name,
        seed=payload.seed,
        is_placeholder=payload.is_placeholder
    )
    db.add(player)
    db.commit()
    db.refresh(player)
    return player


@router.post("/batch", response_model=List[PlayerResponse], status_code=status.HTTP_201_CREATED)
def register_players_batch(
    tournament_id: str,
    payload: PlayerBatchCreate,
    current_user: CurrentUser = Depends(require_organiser),
    db: Session = Depends(get_db)
):
    """Registers multiple players in a single transaction (e.g., from an Excel roster)."""
    tournament = db.query(models.Tournament).filter(models.Tournament.id == tournament_id).first()
    if not tournament:
        raise HTTPException(status_code=404, detail="Tournament not found")
    ensure_tournament_access(tournament, current_user, db)

    if tournament.status != models.TournamentStatus.DRAFT:
        raise HTTPException(
            status_code=400,
            detail="Cannot add players once scheduling has started."
        )

    created_players = []
    for p_in in payload.players:
        player = models.Player(
            id=f"P_{uuid.uuid4().hex[:8]}",
            tournament_id=tournament_id,
            name=p_in.name,
            seed=p_in.seed,
            is_placeholder=p_in.is_placeholder
        )
        db.add(player)
        created_players.append(player)

    db.commit()
    for p in created_players:
        db.refresh(p)
    return created_players


@router.get("", response_model=List[PlayerResponse])
def list_tournament_players(
    tournament_id: str,
    include_placeholders: bool = False,
    current_user: CurrentUser = Depends(require_operator_or_organiser),
    db: Session = Depends(get_db)
):
    """Retrieves all registered players for a tournament."""
    tournament = db.query(models.Tournament).filter(models.Tournament.id == tournament_id).first()
    if not tournament:
        raise HTTPException(status_code=404, detail="Tournament not found")
    ensure_tournament_access(tournament, current_user, db)

    query = db.query(models.Player).filter(models.Player.tournament_id == tournament_id)
    if not include_placeholders:
        query = query.filter(models.Player.is_placeholder == False)

    return query.order_by(models.Player.seed.asc().nullslast()).all()


@router.delete("")
def delete_all_players(
    tournament_id: str,
    current_user: CurrentUser = Depends(require_organiser),
    db: Session = Depends(get_db)
):
    """Removes every player from a DRAFT tournament's roster so a fresh roster can be added."""
    tournament = db.query(models.Tournament).filter(models.Tournament.id == tournament_id).first()
    if not tournament:
        raise HTTPException(status_code=404, detail="Tournament not found")
    ensure_tournament_access(tournament, current_user, db)

    if tournament.status != models.TournamentStatus.DRAFT:
        raise HTTPException(
            status_code=400,
            detail="Cannot remove players once scheduling has started or the tournament is in progress."
        )

    if db.query(models.Match).filter(models.Match.tournament_id == tournament_id).first():
        raise HTTPException(status_code=400, detail="Cannot remove players while fixtures exist for this tournament.")

    deleted = db.query(models.Player).filter(models.Player.tournament_id == tournament_id).delete(synchronize_session=False)
    db.commit()
    return {"message": "Roster cleared.", "deleted_count": deleted}


@router.post("/upload", status_code=status.HTTP_201_CREATED)
async def upload_roster_file(
    tournament_id: str,
    file: UploadFile = File(...),
    current_user: CurrentUser = Depends(require_organiser),
    db: Session = Depends(get_db)
):
    """
    Ingests a CSV or Excel (.xlsx) file of players.
    Columns expected: name (required), email, phone, club_or_city.
    Automatically matches existing athletes or creates new global profiles.
    """
    tournament = db.query(models.Tournament).filter(models.Tournament.id == tournament_id).first()
    if not tournament:
        raise HTTPException(status_code=404, detail="Tournament not found")
    ensure_tournament_access(tournament, current_user, db)

    if tournament.status != models.TournamentStatus.DRAFT:
        raise HTTPException(
            status_code=400,
            detail="Cannot add players once scheduling has started."
        )

    # 1. Read and parse the file
    content = await file.read()
    filename = file.filename.lower()
    
    parsed_rows: List[Dict[str, str]] = []

    try:
        if filename.endswith(".csv"):
            # Parse CSV
            decoded_content = content.decode("utf-8-sig") # utf-8-sig handles Excel CSV exports well
            reader = csv.DictReader(io.StringIO(decoded_content))
            for row in reader:
                parsed_rows.append({k.strip().lower(): v.strip() for k, v in row.items() if k and v})
                
        elif filename.endswith(".xlsx"):
            # Parse Excel
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet = wb.active
            headers = [str(cell.value).strip().lower() for cell in sheet[1] if cell.value]
            
            for row in sheet.iter_rows(min_row=2, values_only=True):
                row_data = {headers[i]: str(row[i]).strip() for i in range(len(headers)) if i < len(row) and row[i] is not None}
                if row_data:
                    parsed_rows.append(row_data)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Please upload a .csv or .xlsx file.")
            
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error parsing file: {str(e)}")

    if not parsed_rows:
        raise HTTPException(status_code=400, detail="The uploaded file contains no data.")

    # 2. Process rows and deduplicate athletes
    stats = {"total_processed": 0, "new_athletes_created": 0, "returning_athletes": 0, "skipped_duplicates": 0}
    
    for row in parsed_rows:
        name = row.get("name")
        if not name:
            continue # Name is mandatory, skip empty/invalid rows
            
        email = row.get("email") or None
        phone = row.get("phone") or None
        club_or_city = row.get("club_or_city") or row.get("club") or row.get("city") or None

        # Look for existing athlete by email, then by phone
        athlete = None
        if email:
            athlete = db.query(models.Athlete).filter(models.Athlete.email == email).first()
        if not athlete and phone:
            athlete = db.query(models.Athlete).filter(models.Athlete.phone == phone).first()

        if athlete:
            stats["returning_athletes"] += 1
            # Update missing info if the returning athlete didn't have it
            if email and not athlete.email: athlete.email = email
            if phone and not athlete.phone: athlete.phone = phone
            if club_or_city and not athlete.club_or_city: athlete.club_or_city = club_or_city
        else:
            # Create new global athlete
            athlete = models.Athlete(
                name=name,
                email=email,
                phone=phone,
                club_or_city=club_or_city
            )
            db.add(athlete)
            db.flush() # Generate ID immediately
            stats["new_athletes_created"] += 1

        # 3. Check if already registered in this specific tournament
        existing_entry = db.query(models.Player).filter(
            models.Player.tournament_id == tournament_id,
            models.Player.athlete_id == athlete.id
        ).first()

        if existing_entry:
            stats["skipped_duplicates"] += 1
            continue

        # 4. Add to Tournament Roster (Notice: No seed is set!)
        player = models.Player(
            id=f"P_{uuid.uuid4().hex[:8]}",
            tournament_id=tournament_id,
            athlete_id=athlete.id,
            name=athlete.name,
            is_placeholder=False
        )
        db.add(player)
        stats["total_processed"] += 1

    db.commit()
    return {
        "message": "Roster upload complete.",
        "stats": stats
    }


# =====================================================================
# WITHDRAWAL (INJURY / EMERGENCY / NO-SHOW)
# =====================================================================
@router.post("/{player_id}/withdraw", response_model=PlayerWithdrawResponse)
def withdraw_player(
    tournament_id: str,
    player_id: str,
    payload: PlayerWithdrawRequest,
    current_user: CurrentUser = Depends(require_organiser),
    db: Session = Depends(get_db),
):
    """
    Marks a registered player as withdrawn. Any of their pending matches whose
    opponent is already known are immediately walked over in the opponent's favor;
    the player is excluded from all future round pairings and knockout qualification.
    Already-played results are untouched.
    """
    tournament = db.query(models.Tournament).filter(models.Tournament.id == tournament_id).first()
    if not tournament:
        raise HTTPException(status_code=404, detail="Tournament not found")
    ensure_tournament_access(tournament, current_user, db)

    player = db.query(models.Player).filter(
        models.Player.id == player_id,
        models.Player.tournament_id == tournament_id,
    ).first()
    if not player:
        raise HTTPException(status_code=404, detail="Player not found in this tournament.")
    if player.is_placeholder:
        raise HTTPException(status_code=400, detail="Cannot withdraw a TBD placeholder slot.")
    if player.is_withdrawn:
        raise HTTPException(status_code=400, detail="Player has already withdrawn.")

    player.is_withdrawn = True
    player.withdrawn_at = datetime.utcnow()
    player.withdrawal_reason = payload.reason
    db.flush()

    walkover_matches = resolve_pending_walkovers(tournament_id, db)

    db.commit()
    db.refresh(player)
    for m in walkover_matches:
        db.refresh(m)

    return PlayerWithdrawResponse(player=player, walkover_matches=walkover_matches)